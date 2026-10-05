#!/usr/bin/env python3
"""Data rating for every position tab, from the StatsBomb player-season exports in data/ (gitignored).

For each position:
  1. Season rating (0-100) = weighted family scores; a family is the mean percentile of its metrics vs players
     in the same position pool, league and season. Families and weights come from that tab's "WHAT WE NEED" list.
  2. Level-adjusted = 0.6 * rating + 40 * league strength, so leagues sit on one scale.
  3. Player rating = seasons weighted by recency x minutes reliability (mins / 900, max 1).

Usage:  python3 position_rating.py CB          # top 25 + Bristol benchmark
        python3 position_rating.py Winger --all
"""
import argparse, os, re, unicodedata
from datetime import date
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "data")
TODAY = date(2026, 10, 5)

STRENGTH = {"Championship": 1.0, "League One": 0.8, "Scotland Premiership": 0.8, "League Two": 0.6,
            "National League": 0.35, "K League 1": 0.8}  # K League 1 placed with L1/SPL - an assumption
SEASON_W = {"2026/2027": 0.30, "2025/2026": 0.45, "2024/2025": 0.25, "2025": 0.45, "2026": 0.30}

EN5 = "positions/PlayerSeason_{}_EN5_300mins.xlsx"
# tab: files, position filter (None = all rows), {family: (label, weight, [(metric, higher_is_better)])}
POSITIONS = {
    "GK": ([EN5.format("GK")], None, {
        "shotstop": ("Shot-stopping", 35, [("gsaa_90", True), ("save_ratio", True), ("obv_gk_90", True)]),
        "crosses":  ("Claims crosses", 20, [("clcaa", True)]),
        "distrib":  ("Calm distribution", 25, [("pressured_passing_ratio", True), ("passing_ratio", True),
                                                ("errors_90", False)]),
        "sweep":    ("Sweeps", 20, [("da_aggressive_distance", True), ("average_x_defensive_action", True)]),
    }),
    "CB": (["CB_all_leagues_3seasons.xlsx"], None, {
        "aerial":   ("Aerial", 20, [("aerial_ratio", True), ("aerial_wins_90", True)]),
        "secure":   ("Secure under press", 20, [("pressured_passing_ratio", True), ("errors_90", False),
                                                ("turnovers_90", False), ("dispossessions_90", False)]),
        "winback":  ("Wins it back", 15, [("padj_tackles_and_interceptions_90", True), ("ball_recoveries_90", True),
                                          ("defensive_action_regains_90", True)]),
        "1v1":      ("1v1", 15, [("challenge_ratio", True), ("dribbled_past_90", False)]),
        "clear":    ("Clearances/blocks", 10, [("padj_clearances_90", True), ("blocks_per_shot", True)]),
        "setpiece": ("Set-piece goals", 10, [("npg_90", True), ("np_xg_90", True)]),
        "press":    ("Pressure regains", 5, [("pressure_regains_90", True)]),
        "value":    ("Defensive OBV", 5, [("obv_defensive_action_90", True)]),
    }),
    "LB": ([EN5.format("FB")], {"Left Back", "Left Wing Back"}, {
        "supply":   ("Box supply", 30, [("op_passes_into_box_90", True), ("op_xa_90", True),
                                        ("op_key_passes_90", True)]),
        "carry":    ("Carries up the pitch", 25, [("carries_90", True), ("deep_progressions_90", True),
                                                  ("obv_dribble_carry_90", True)]),
        "1v1":      ("1v1 defending", 25, [("challenge_ratio", True), ("dribbled_past_90", False),
                                           ("padj_tackles_90", True)]),
        "setpiece": ("Set-piece delivery", 10, [("sp_key_passes_90", True), ("sp_xa_90", True)]),
        "secure":   ("Secure on the ball", 10, [("turnovers_90", False), ("pressured_passing_ratio", True)]),
    }),
    "RB": ([EN5.format("FB")], {"Right Back", "Right Wing Back"}, {
        "defend":   ("Defends first", 30, [("padj_tackles_and_interceptions_90", True),
                                           ("defensive_action_regains_90", True), ("padj_clearances_90", True)]),
        "supply":   ("Box supply", 20, [("op_passes_into_box_90", True), ("op_xa_90", True),
                                        ("op_key_passes_90", True)]),
        "intensity": ("High-intensity work", 15, [("pressures_90", True), ("counterpressures_90", True)]),
        "1v1":      ("1v1 defending", 20, [("challenge_ratio", True), ("dribbled_past_90", False)]),
        "cross":    ("Crossing quality", 15, [("crossing_ratio", True), ("crosses_90", True)]),
    }),
    "CM 6": ([EN5.format("DMCM")], {"Left Defensive Midfielder", "Right Defensive Midfielder",
                                    "Centre Defensive Midfielder"}, {
        "screen":   ("Screens / clears", 20, [("aerial_wins_90", True), ("padj_clearances_90", True),
                                              ("padj_interceptions_90", True)]),
        "winback":  ("Wins it back", 25, [("padj_tackles_and_interceptions_90", True), ("ball_recoveries_90", True),
                                          ("defensive_action_regains_90", True)]),
        "resist":   ("Press-resistant", 25, [("pressured_passing_ratio", True), ("turnovers_90", False),
                                             ("dispossessions_90", False)]),
        "transition": ("Covers ground in transition", 15, [("counterpressures_90", True),
                                                          ("counterpressure_regains_90", True)]),
        "composure": ("Composure in possession", 15, [("passing_ratio", True), ("obv_pass_90", True)]),
    }),
    "CM 8": ([EN5.format("DMCM")], None, {
        "supply":   ("Box supply", 20, [("op_passes_into_box_90", True), ("deep_completions_90", True),
                                        ("op_xa_90", True)]),
        "winback":  ("Wins it back / counter-press", 20, [("padj_tackles_and_interceptions_90", True),
                                                          ("counterpressures_90", True),
                                                          ("counterpressure_regains_90", True)]),
        "progress": ("Progresses the ball", 25, [("deep_progressions_90", True), ("carries_90", True),
                                                 ("obv_pass_90", True), ("obv_dribble_carry_90", True)]),
        "box":      ("Arrives in the box", 20, [("np_shots_90", True), ("np_xg_90", True),
                                                ("touches_inside_box_90", True)]),
        "secure":   ("Secure under pressure", 15, [("pressured_passing_ratio", True), ("turnovers_90", False)]),
    }),
    "CM 10": ([EN5.format("AM")], {"Centre Attacking Midfielder"}, {
        "box":      ("Gets us into the box", 25, [("touches_inside_box_90", True), ("op_passes_into_box_90", True)]),
        "create":   ("Final-third creator", 30, [("op_key_passes_90", True), ("op_xa_90", True)]),
        "setpiece": ("Set-piece taker", 10, [("sp_key_passes_90", True), ("sp_xa_90", True)]),
        "secure":   ("Keeps the ball under pressure", 20, [("pressured_passing_ratio", True),
                                                          ("turnovers_90", False), ("dispossessions_90", False)]),
        "press":    ("Presses from the front", 15, [("pressures_90", True), ("pressure_regains_90", True)]),
    }),
    # English wingers: only the wide AMs in the AM file so far - an EN5 wingers export is still needed
    "Winger": ([EN5.format("AM"), "positions/Wingers_K_League_1_2025.xlsx"],
               {"Left Attacking Midfielder", "Right Attacking Midfielder", "Left Wing", "Right Wing",
                "Left Midfielder", "Right Midfielder"}, {
        "box":      ("Gets into the box", 30, [("np_shots_90", True), ("touches_inside_box_90", True),
                                               ("np_xg_90", True)]),
        "1v1":      ("Beats a man", 25, [("dribble_ratio", True), ("dribbles_90", True)]),
        "transition": ("Transition threat", 15, [("obv_dribble_carry_90", True),
                                                 ("f3_lbp_to_space_10_received_90", True)]),
        "retain":   ("Keeps the ball", 20, [("dispossessions_90", False), ("turnovers_90", False)]),
        "workback": ("Works back", 10, [("padj_tackles_and_interceptions_90", True), ("pressure_regains_90", True)]),
    }),
    # Target / poacher: knows where the goal is when supplied; aerial still matters, pressing least
    "CF Target": ([EN5.format("CF")], None, {
        "finish":   ("Finishing", 40, [("np_xg_90", True), ("npg_90", True), ("np_xg_per_shot", True),
                                       ("conversion_ratio", True)]),
        "box":      ("Box presence", 25, [("touches_inside_box_90", True), ("np_shots_90", True)]),
        "aerial":   ("Aerial / first contacts", 20, [("aerial_wins_90", True), ("aerial_ratio", True)]),
        "holdup":   ("Holds it up", 10, [("fouls_won_90", True), ("dispossessions_90", False), ("turnovers_90", False)]),
        "press":    ("Presses from the front", 5, [("pressures_90", True), ("pressure_regains_90", True)]),
    }),
    "CF Runner": ([EN5.format("CF")], None, {
        "behind":   ("Runs in behind", 25, [("f3_lbp_to_space_10_received_90", True), ("obv_dribble_carry_90", True)]),
        "finish":   ("Box presence & finishing", 30, [("np_xg_90", True), ("npg_90", True), ("np_xg_per_shot", True)]),
        "press":    ("Presses from the front", 20, [("pressures_90", True), ("pressure_regains_90", True)]),
        "link":     ("Links play", 25, [("op_key_passes_90", True), ("op_xa_90", True), ("op_xgbuildup_90", True)]),
    }),
}


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z ]", "", s.replace("-", " ")).split()


def load(tab):
    files, positions, _ = POSITIONS[tab]
    rows, seen = [], set()
    for f in files:
        wb = load_workbook(os.path.join(DATA, f), read_only=True, data_only=True)
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            head = [h.replace("player_season_", "") for h in next(it)]
            for r in it:
                if not r[0]: continue
                d = dict(zip(head, r))
                if positions and d.get("primary_position") not in positions: continue
                d["league"] = d.get("league") or d.get("competition_name")
                key = (d["player_id"], d["season_name"], d["team_id"])
                if key in seen or d["league"] not in STRENGTH: continue  # tabs repeat rows ("All" + per league)
                seen.add(key)
                d["name"] = d.get("player_known_name") or d["player_name"]
                rows.append(d)
    return rows


def pct(v, vals, hi):
    if not isinstance(v, (int, float)): return None
    vals = [x for x in vals if isinstance(x, (int, float))]
    better = sum(1 for x in vals if (x < v if hi else x > v))
    return 100 * (better + 0.5 * sum(1 for x in vals if x == v)) / len(vals)


def rate(tab, rows):
    fams = POSITIONS[tab][2]
    groups = {}
    for r in rows: groups.setdefault((r["league"], r["season_name"]), []).append(r)
    for peers in groups.values():
        cols = {c: [p.get(c) for p in peers] for _, _, ms in fams.values() for c, _ in ms}
        for r in peers:
            fam, tot, wsum = {}, 0, 0
            for f, (_, w, ms) in fams.items():
                ps = [p for p in (pct(r.get(c), cols[c], hi) for c, hi in ms) if p is not None]
                if not ps: continue
                fam[f] = sum(ps) / len(ps); tot += w * fam[f]; wsum += w
            r["fam"], r["rating"] = fam, (tot / wsum if wsum else None)
            r["adj"] = 0.6 * r["rating"] + 40 * STRENGTH[r["league"]] if r["rating"] is not None else None
    return rows


def players(tab, rows):
    fams = POSITIONS[tab][2]
    by = {}
    for r in rows: by.setdefault(r["player_id"], []).append(r)
    out = []
    for pid, rs in by.items():
        rs = [r for r in rs if r["adj"] is not None]
        if not rs: continue
        w = [(SEASON_W.get(r["season_name"], 0.25) * min(r["minutes"] / 900, 1), r) for r in rs]
        tw = sum(x for x, _ in w)
        latest = max(rs, key=lambda r: (r["season_name"], r["minutes"]))
        fam = {}
        for f in fams:
            vals = [(x, r["fam"][f]) for x, r in w if f in r["fam"]]
            if vals: fam[f] = sum(x * v for x, v in vals) / sum(x for x, _ in vals)
        by_season = {}
        for r in rs:  # a mid-season move gives two rows; keep the bigger one per season
            if r["season_name"] not in by_season or r["minutes"] > by_season[r["season_name"]]["minutes"]:
                by_season[r["season_name"]] = r
        seasons = sorted(by_season)
        trend = (by_season[seasons[-1]]["adj"] - by_season[seasons[-2]]["adj"]) if len(seasons) > 1 else None
        age, bd = None, latest.get("birth_date")
        if bd:
            y, m, d = map(int, str(bd)[:10].split("-"))
            age = TODAY.year - y - ((TODAY.month, TODAY.day) < (m, d))
        lfr = sum(r["minutes"] * (r.get("left_foot_ratio") or 0) for r in rs) / sum(r["minutes"] for r in rs)
        out.append({
            "id": pid, "name": latest["name"], "team": latest["team_name"], "league": latest["league"],
            "season": latest["season_name"], "age": age, "height": latest.get("player_height"),
            "foot": "L" if lfr > 0.5 else "R", "position": latest.get("primary_position"),
            "mins_now": round(sum(r["minutes"] for r in rs if r["season_name"] in ("2026/2027", "2026"))),
            "rating": sum(x * r["adj"] for x, r in w) / tw, "fam": fam, "trend": trend,
            "history": "; ".join(f"{s} {by_season[s]['team_name']} ({by_season[s]['league']}, "
                                 f"{round(by_season[s]['minutes'])}') {by_season[s]['adj']:.0f}" for s in seasons),
        })
    return out


def rated(tab):
    return players(tab, rate(tab, load(tab)))


def strengths(tab, p, n=3):
    fams = POSITIONS[tab][2]
    top = sorted(p["fam"].items(), key=lambda kv: -kv[1])
    return ", ".join(fams[f][0] for f, v in top[:n] if v >= 65) or "-"


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("tab", choices=list(POSITIONS))
    ap.add_argument("--all", action="store_true")
    a = ap.parse_args()
    ps = sorted(rated(a.tab), key=lambda p: -p["rating"])
    for p in ps if a.all else ps[:25]:
        print(f"{p['rating']:5.1f} {p['name'][:26]:26} {p['team'][:22]:22} {p['league'][:20]:20} age {p['age']} "
              f"{p['foot']} now {p['mins_now']}' | {strengths(a.tab, p)}")
    print("\nBristol Rovers:")
    for p in ps:
        if "Bristol Rovers" in p["history"]:
            print(f"{p['rating']:5.1f} {p['name']} | {p['history']}")
